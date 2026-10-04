# -*- coding: utf-8 -*-
"""build_docx.py — 把审计报告 Markdown 渲染为 DOCX（中文排版）。

用法：
    python build_docx.py <input.md> <output.docx> [标题]

支持：# / ## / ### 标题、段落、- 列表、有序列表、> 引用、--- 分隔线、
| 表格 |、**粗体** 行内标记、`代码` 行内标记。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT_EA = "微软雅黑"
BODY_FONT_LATIN = "Segoe UI"
MONO_FONT = "Consolas"


def _set_run_font(run, name_latin: str, name_ea: str, size: float | None = None,
                  bold: bool | None = None, color: tuple[int, int, int] | None = None) -> None:
    run.font.name = name_latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name_latin)
    rfonts.set(qn("w:hAnsi"), name_latin)
    rfonts.set(qn("w:eastAsia"), name_ea)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(paragraph, text: str, base_size: float = 10.5,
             base_color: tuple[int, int, int] | None = None) -> None:
    """把含 **粗体** 与 `代码` 的文本写入段落。"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            # 粗体内可能内嵌行内代码标记（**`x`**），去掉反引号避免字面显示
            run = paragraph.add_run(part[2:-2].replace("`", ""))
            _set_run_font(run, BODY_FONT_LATIN, BODY_FONT_EA, base_size, True, base_color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, MONO_FONT, MONO_FONT, base_size - 0.5, False, (0x9B, 0x2C, 0x2C))
        else:
            run = paragraph.add_run(part)
            _set_run_font(run, BODY_FONT_LATIN, BODY_FONT_EA, base_size, False, base_color)


def setup_document(doc: Document, title: str) -> None:
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = BODY_FONT_LATIN
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_EA)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.18

    for level, size in ((1, 17), (2, 14), (3, 12), (4, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x11, 0x33, 0x55)
        style.font.name = BODY_FONT_LATIN
        style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_EA)

    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # 页眉
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = ""
    hrun = hp.add_run(title)
    _set_run_font(hrun, BODY_FONT_LATIN, BODY_FONT_EA, 8.5, False, (0x66, 0x66, 0x66))
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # 页脚页码
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    _set_run_font(run, BODY_FONT_LATIN, BODY_FONT_EA, 8.5, False, (0x66, 0x66, 0x66))
    fld_begin = run._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr = run._element.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
    instr.text = " PAGE "
    fld_end = run._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    run._element.append(fld_begin)
    run._element.append(instr)
    run._element.append(fld_end)


def split_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def is_separator_row(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip())) and "-" in line


def add_table(doc: Document, rows: list[list[str]]) -> None:
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(ncols):
            text = row[j] if j < len(row) else ""
            cell = cells[j]
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(1)
            add_runs(para, text, base_size=8.5, base_color=None)
            if i == 0:
                for run in para.runs:
                    run.font.bold = True
    doc.add_paragraph()


def convert(md_path: Path, out_path: Path, title: str) -> int:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup_document(doc, title)

    i = 0
    n_tables = 0
    first_heading_done = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator_row(lines[i + 1]):
            rows = [split_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            n_tables += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", stripped):
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            run = para.add_run("─" * 40)
            _set_run_font(run, BODY_FONT_LATIN, BODY_FONT_EA, 8, False, (0xBB, 0xBB, 0xBB))
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            text = m.group(2).strip()
            if not first_heading_done and level == 1:
                first_heading_done = True
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(text)
                _set_run_font(run, BODY_FONT_LATIN, BODY_FONT_EA, 22, True, (0x0B, 0x2A, 0x4A))
                para.paragraph_format.space_after = Pt(14)
                i += 1
                continue
            heading = doc.add_heading(level=level)
            base = {1: 17, 2: 14, 3: 12, 4: 11}[level]
            add_runs(heading, text, base_size=base)
            for run in heading.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x11, 0x33, 0x55)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            add_runs(para, text, base_size=9.5, base_color=(0x44, 0x44, 0x44))
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            para = doc.add_paragraph(style="List Bullet")
            add_runs(para, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if m:
            para = doc.add_paragraph(style="List Number")
            add_runs(para, m.group(1))
            i += 1
            continue

        # 普通段落
        para = doc.add_paragraph()
        add_runs(para, stripped)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return n_tables


def main() -> int:
    if len(sys.argv) < 3:
        print("usage: build_docx.py <input.md> <output.docx> [title]")
        return 2
    md = Path(sys.argv[1])
    out = Path(sys.argv[2])
    title = sys.argv[3] if len(sys.argv) > 3 else "鲲鹏量化交易系统审计报告"
    tables = convert(md, out, title)
    print(f"OK {out} tables={tables} bytes={out.stat().st_size}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
